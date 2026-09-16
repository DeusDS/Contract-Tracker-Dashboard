"""Publish an Aging Project Tracker Excel refresh to the GitHub Pages dashboard.

Run with no arguments to select an Excel export through Windows Explorer. The
script updates the one hosted dashboard, commits the change, and pushes main.
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import re
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

REPO = Path(r"C:\Users\damon\OneDrive - Atlanta Regional Commission\ARC\Desktop\GitHub\Aging Dashboards\Aging Contract Dashboards\Contract Tracker Dashboard")
DEPARTMENT = "Aging & Independence Services"
STAGES = [
    ("Procurement Request Status", "Procurement"),
    ("Sole Source Request Status", "Sole Source"),
    ("Vendor Selection Status", "Vendor Selection"),
    ("Amendment Request Status", "Amendment"),
    ("Contract Request Status", "Contract Request"),
    ("Contract Execution Status", "Contract Execution"),
]
DATE_RE = re.compile(r"\b\d{1,2}/\d{1,2}/\d{4}\b")


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<\s*/?br\s*/?\s*>", " | ", str("" if value is None else value), flags=re.I))).strip(" | ")


def parse_status(value: object) -> tuple[str, str, str]:
    text = clean(value)
    match = DATE_RE.search(text)
    if match:
        return text[:match.start()].strip(" | "), match.group(), text[match.end():].strip(" | ")
    parts = [part.strip() for part in text.split("|")]
    return parts[0] if parts else "", "", " | ".join(parts[1:]).strip()


def is_terminal(status: str) -> bool:
    normalized = clean(status).casefold()
    if "partial" in normalized:
        return False
    return normalized in {"complete", "completed", "approvals completed", "executed", "fully executed"}


def parse_days(value: object) -> int | float | str:
    text = clean(value)
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError:
        return ""
    return int(number) if number.is_integer() else number


def age_from_date(value: object, as_of: date) -> int | str:
    text = clean(value)
    if not text:
        return ""
    try:
        action_date = datetime.strptime(text, "%m/%d/%Y").date()
    except ValueError:
        return ""
    return max(0, (as_of - action_date).days)


def ensure_project_age(project: dict[str, object], as_of: date) -> dict[str, object]:
    source_age = parse_days(project.get("days_in_current_stage"))
    if source_age != "":
        project["days_in_current_stage"] = source_age
        project["age_source"] = clean(project.get("age_source")) or "Source export"
        return project
    derived_age = age_from_date(project.get("current_stage_date"), as_of)
    project["days_in_current_stage"] = derived_age
    project["age_source"] = "Derived from current workflow date" if derived_age != "" else ""
    return project


def sheet_records(path: Path, name: str) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[name]
    headers = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    return [dict(zip(headers, row)) for row in worksheet.iter_rows(min_row=2, values_only=True)]


def csv_records(path: Path) -> list[dict[str, object]]:
    """Find the real header row and read CSV exports with leading blank lines."""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    required = {"Requisition Id", "Contract Request Status"}
    for index, row in enumerate(rows[:20]):
        headers = [clean(value) for value in row]
        if required.issubset(headers):
            return [
                dict(zip(headers, values))
                for values in rows[index + 1:]
                if any(clean(value) for value in values)
            ]
    return []


def raw_records(path: Path) -> list[dict[str, object]]:
    if path.suffix.lower() == ".csv":
        records = csv_records(path)
        required = {"Requisition Id", "Contract Request Status"}
        if not records or not required.issubset(records[0]):
            raise ValueError("The CSV does not contain a recognizable Project Tracker export.")
        return records
    workbook = load_workbook(path, read_only=True, data_only=True)
    matches: list[tuple[int, list[object], list[tuple[object, ...]]]] = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        for index, row in enumerate(rows[:6]):
            if "Requisition Id" in row and "Contract Request Status" in row:
                matches.append((len(rows) - index - 1, list(row), rows[index + 1:]))
    if not matches:
        raise ValueError("The workbook does not contain a recognizable Project Tracker export.")
    _, headers, values = max(matches, key=lambda item: item[0])
    return [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]


def normalize_raw(path: Path) -> dict[str, object]:
    projects: list[dict[str, object]] = []
    events: list[dict[str, object]] = []
    issues: list[dict[str, object]] = []
    data_as_of = datetime.fromtimestamp(path.stat().st_mtime).date()
    for source_row in raw_records(path):
        if clean(source_row.get("Department")) != DEPARTMENT:
            continue
        requisition_id = clean(source_row.get("Requisition Id"))
        key = requisition_id or f"row-{len(projects) + 1}"
        project_events: list[dict[str, object]] = []
        for rank, (column, portal) in enumerate(STAGES, 1):
            status, action_date, waiting_on = parse_status(source_row.get(column))
            if status:
                project_events.append({"project_key": key, "stage_rank": rank, "portal": portal, "status": status, "action_date": action_date, "waiting_on": waiting_on, "source": "Derived from loaded export"})
        current = (next((event for event in reversed(project_events) if not is_terminal(str(event["status"]))), None) or (project_events[-1] if project_events else {}))
        current_portal = str(current.get("portal", "Unclassified"))
        current_date = clean(current.get("action_date"))
        date_source_portal = current_portal if current_date else ""
        if not current_date and current:
            eligible_dates = [
                event for event in project_events
                if clean(event.get("action_date")) and int(event.get("stage_rank", 0)) <= int(current.get("stage_rank", 0))
            ]
            if eligible_dates:
                dated_event = max(eligible_dates, key=lambda event: datetime.strptime(str(event["action_date"]), "%m/%d/%Y"))
                current_date = str(dated_event["action_date"])
                date_source_portal = str(dated_event["portal"])
        days = parse_days(source_row.get("Days Since Contract Request")) if current_portal == "Contract Request" else ""
        project = {
            "project_key": key,
            "requisition_id": requisition_id,
            "contract_number": "",
            "project_name": clean(source_row.get("Project Name")),
            "vendor": clean(source_row.get("Vendor")),
            "department": DEPARTMENT,
            "project_manager": clean(source_row.get("Project Manager")),
            "procurement_type": clean(source_row.get("Procurement Type")),
            "document_type": clean(source_row.get("Document Type")),
            "record_state": "Active (source export)",
            "current_portal": current_portal,
            "current_status": str(current.get("status", "No workflow status")),
            "current_stage_date": current_date,
            "stage_date_source": "Current workflow date" if current_date and date_source_portal == current_portal else (f"Latest known date from {date_source_portal}" if current_date else ""),
            "days_in_current_stage": days,
            "age_source": "Source: Days Since Contract Request" if days != "" else "",
            "waiting_on": str(current.get("waiting_on", "")),
            "stage_source": "Derived from loaded export",
        }
        projects.append(ensure_project_age(project, data_as_of))
        events.extend(project_events)
        for field in ("project_name", "requisition_id", "contract_number", "project_manager"):
            if not str(project[field]).strip():
                issues.append({"project_key": key, "project_name": project["project_name"], "field": field.replace("_", " ").title(), "issue": "Missing from loaded export", "severity": "High" if field in {"project_name", "requisition_id"} else "Medium"})
    return {"projects": projects, "events": events, "issues": issues}


def payload_from_workbook(path: Path) -> dict[str, object]:
    if path.suffix.lower() == ".csv":
        data = normalize_raw(path)
        return add_metadata(data, path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Projects" in workbook.sheetnames:
        projects = [row for row in sheet_records(path, "Projects") if clean(row.get("department")) == DEPARTMENT]
        workbook_as_of = datetime.fromtimestamp(path.stat().st_mtime).date()
        for project in projects:
            if not clean(project.get("project_manager")):
                project["project_manager"] = clean(project.get("program_manager"))
            project.pop("program_manager", None)
            ensure_project_age(project, workbook_as_of)
        keys = {row.get("project_key") for row in projects}
        events = [row for row in sheet_records(path, "Workflow Events") if row.get("project_key") in keys] if "Workflow Events" in workbook.sheetnames else []
        issues = [row for row in sheet_records(path, "Data Quality") if row.get("project_key") in keys] if "Data Quality" in workbook.sheetnames else []
        data = {"projects": projects, "events": events, "issues": issues}
    else:
        data = normalize_raw(path)
    return add_metadata(data, path)


def add_metadata(data: dict[str, object], path: Path) -> dict[str, object]:
    if not data["projects"]:
        raise ValueError(f"No {DEPARTMENT} project records were found in the selected workbook.")
    data["meta"] = {
        "source_file": path.name,
        "data_as_of": datetime.fromtimestamp(path.stat().st_mtime).date().isoformat(),
        "department": DEPARTMENT,
        "records": len(data["projects"]),
        "events": len(data["events"]),
        "issues": len(data["issues"]),
        "note": "Published from the latest Excel or CSV export by the local dashboard publisher.",
    }
    return data


def hosted_dashboard() -> Path:
    """Find the active GitHub Pages document without relying on historical names."""
    index = REPO / "index.html"
    if index.exists():
        return index
    candidates = [path for path in REPO.glob("*.html") if "const INITIAL=" in path.read_text(encoding="utf-8", errors="ignore")]
    if len(candidates) == 1:
        return candidates[0]
    raise FileNotFoundError("Could not identify the hosted dashboard HTML file. Expected index.html or one HTML file containing the dashboard data block.")


def git(*args: str) -> str:
    result = subprocess.run(["git", "-C", str(REPO), *args], capture_output=True, text=True)
    if result.returncode:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip())
    return result.stdout.strip()


def commit_and_push(dashboard: Path, message: str) -> bool:
    """Commit changed dashboard data when needed, then publish pending commits."""
    git("add", "--", dashboard.name)
    staged_files = git("diff", "--cached", "--name-only")
    committed = bool(staged_files)
    if committed:
        git("commit", "-m", message)
    else:
        print("Dashboard data is unchanged; no new commit was needed.")
    git("push", "origin", "main")
    return committed


def pick_workbook() -> Path:
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk(); root.withdraw(); root.attributes("-topmost", True)
    selected = filedialog.askopenfilename(title="Select the latest Project Tracker Excel or CSV export", filetypes=[("Excel or CSV", "*.xlsx *.xls *.csv"), ("Excel workbooks", "*.xlsx *.xls"), ("CSV files", "*.csv"), ("All files", "*.*")])
    root.destroy()
    if not selected:
        raise SystemExit("No workbook selected. Nothing was published.")
    return Path(selected)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, help="Excel or CSV export to publish; omit to open the Windows file picker.")
    parser.add_argument("--dry-run", action="store_true", help="Validate the selected Excel or CSV export without changing or publishing the dashboard.")
    args = parser.parse_args()
    source = args.source or pick_workbook()
    if not source.exists():
        raise FileNotFoundError(source)
    payload = payload_from_workbook(source)
    if args.dry_run:
        print(f"Validated {payload['meta']['records']} Aging projects, {payload['meta']['events']} events, and {payload['meta']['issues']} exceptions.")
        return
    dashboard = hosted_dashboard()
    if git("status", "--porcelain"):
        raise RuntimeError("Repository has uncommitted changes. Commit or stash them before publishing a dashboard refresh.")
    document = dashboard.read_text(encoding="utf-8")
    replacement = "const INITIAL=" + json.dumps(payload, ensure_ascii=False).replace("</", "<\\/") + ";let data="
    updated, count = re.subn(r"const INITIAL=.*?;let data=", replacement, document, count=1, flags=re.DOTALL)
    if count != 1:
        raise RuntimeError("Could not locate the embedded dashboard data block.")
    dashboard.write_text(updated, encoding="utf-8")
    commit_and_push(
        dashboard,
        f"Refresh Aging contract tracker data ({payload['meta']['data_as_of']})",
    )
    print(f"Published {payload['meta']['records']} Aging projects to https://deusds.github.io/Contract-Tracker-Dashboard/")


if __name__ == "__main__":
    main()
